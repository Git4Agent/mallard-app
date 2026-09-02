from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    "docs/MAINLAND_CHINA_OBJECT_STORAGE_COMPATIBILITY_ASSESSMENT_2026-09-01.docx"
)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
INK = "20262E"
MUTED = "5A6573"
LIGHT_GRAY = "F2F4F7"
GRID = "CAD2DC"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GREEN = "2F6B4F"
AMBER = "7A5A00"
RED = "9B1C1C"


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), name)
    r_pr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=120):
    if sum(widths_dxa) != 9360:
        raise ValueError(f"Table widths must total 9360 DXA: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
    set_table_borders(table)


def repeat_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def style_table(table, *, body_size=8.6, status_column: int | None = None):
    repeat_header_row(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            if row_index == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                if status_column is not None and cell_index == status_column:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=body_size,
                        color=NAVY if row_index == 0 else INK,
                        bold=row_index == 0,
                    )


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    r_pr.append(size)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def paragraph_border_bottom(paragraph, color=BLUE, size="12", space="5"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def create_numbering_id(
    doc,
    *,
    number_format: str,
    level_text_value: str,
    font: str | None = None,
):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), number_format)
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), level_text_value)
    level.append(level_text)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)
    if font:
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), font)
        r_fonts.set(qn("w:hAnsi"), font)
        r_pr.append(r_fonts)
        level.append(r_pr)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.insert(0, num_pr)


def add_list(doc, items, *, ordered=False):
    if ordered:
        num_id = create_numbering_id(
            doc,
            number_format="decimal",
            level_text_value="%1.",
        )
    else:
        num_id = create_numbering_id(
            doc,
            number_format="bullet",
            level_text_value="\u2022",
            font="Arial",
        )
    for item in items:
        if isinstance(item, tuple):
            label, text = item
        else:
            label, text = None, item
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.167
        paragraph.paragraph_format.keep_together = True
        apply_numbering(paragraph, num_id)
        if label:
            lead = paragraph.add_run(f"{label}: ")
            set_run_font(lead, bold=True, color=NAVY)
        body = paragraph.add_run(text)
        set_run_font(body, color=INK)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    title_p_pr = title._element.get_or_add_pPr()
    inherited_border = title_p_pr.find(qn("w:pBdr"))
    if inherited_border is not None:
        title_p_pr.remove(inherited_border)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.keep_with_next = True

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    source = doc.styles.add_style("Source Register", 1)
    source.font.name = "Calibri"
    source._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    source._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    source.font.size = Pt(9.5)
    source.font.color.rgb = RGBColor.from_string(INK)
    source.paragraph_format.space_before = Pt(0)
    source.paragraph_format.space_after = Pt(4)
    source.paragraph_format.line_spacing = 1.05
    source.paragraph_format.left_indent = Inches(0.25)
    source.paragraph_format.first_line_indent = Inches(-0.25)
    source.paragraph_format.keep_together = True


def configure_page(doc):
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    for header in (section.header, section.even_page_header):
        paragraph = header.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
        left = paragraph.add_run("MALLARD | TECHNICAL RESEARCH")
        set_run_font(left, size=8.5, color=MUTED, bold=True)
        paragraph.add_run("\t")
        right = paragraph.add_run("MAINLAND CHINA OBJECT STORAGE")
        set_run_font(right, size=8.5, color=MUTED, bold=True)

    for footer in (section.footer, section.even_page_footer):
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        prefix = paragraph.add_run("Page ")
        set_run_font(prefix, size=9, color=MUTED)
        add_field(paragraph, " PAGE ")
        middle = paragraph.add_run(" of ")
        set_run_font(middle, size=9, color=MUTED)
        add_field(paragraph, " NUMPAGES ")


def add_metadata_row(doc, label: str, value: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    lead = paragraph.add_run(f"{label}: ")
    set_run_font(lead, size=10.5, color=NAVY, bold=True)
    body = paragraph.add_run(value)
    set_run_font(body, size=10.5, color=INK)


def add_callout(doc, label: str, text: str, tone="neutral"):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.line_spacing = 1.10
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CALLOUT)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left_border = OxmlElement("w:left")
    color = {
        "positive": GREEN,
        "caution": AMBER,
        "risk": RED,
    }.get(tone, BLUE)
    left_border.set(qn("w:val"), "single")
    left_border.set(qn("w:sz"), "20")
    left_border.set(qn("w:space"), "8")
    left_border.set(qn("w:color"), color)
    borders.append(left_border)
    p_pr.append(borders)
    lead = paragraph.add_run(f"{label}: ")
    set_run_font(lead, color=color, bold=True)
    body = paragraph.add_run(text)
    set_run_font(body, color=INK)


def add_body(doc, text: str, *, italic=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, color=INK, italic=italic)
    return paragraph


def add_table_source(doc, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, size=8.5, color=MUTED, italic=True)


def add_source(doc, source_id: str, title: str, url: str | None, note: str):
    paragraph = doc.add_paragraph(style="Source Register")
    lead = paragraph.add_run(f"[{source_id}] ")
    set_run_font(lead, size=9.5, color=NAVY, bold=True)
    if url:
        add_hyperlink(paragraph, title, url)
    else:
        title_run = paragraph.add_run(title)
        set_run_font(title_run, size=9.5, color=INK)
    note_run = paragraph.add_run(f" - {note}")
    set_run_font(note_run, size=9.5, color=MUTED)


def add_provider_heading(doc, name: str, status: str, page_break_before: bool = False):
    heading = doc.add_heading(name, level=3)
    heading.paragraph_format.page_break_before = page_break_before
    heading.paragraph_format.keep_with_next = True
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(status)
    set_run_font(run, size=10.5, color=NAVY, bold=True)


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    props = doc.core_properties
    props.title = "Mainland China Object Storage Compatibility Assessment for Mallard"
    props.subject = "Provider research and implementation feasibility"
    props.author = ""
    props.last_modified_by = ""
    props.keywords = "Mallard, China, object storage, S3, OSS, COS, OBS, BOS, TOS"

    title = doc.add_paragraph(style="Title")
    title.add_run("Mainland China Object Storage Compatibility Assessment")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Provider research and implementation feasibility for Mallard")
    add_metadata_row(doc, "Prepared for", "Mallard maintainers")
    add_metadata_row(doc, "Research date", "1 September 2026")
    add_metadata_row(doc, "Scope", "Mainland China cloud leaders and prominent object-storage specialists")
    add_metadata_row(doc, "Method", "Official provider documentation, current market research, and repository inspection")
    add_metadata_row(doc, "Change status", "Research only; no product, schema, or safety change")
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(4)
    rule.paragraph_format.space_after = Pt(10)
    paragraph_border_bottom(rule)

    doc.add_heading("Executive summary", level=1)
    add_callout(
        doc,
        "Overall result",
        "Mallard supports none of the reviewed mainland China providers through its current product path. The frontend and backend accept only canonical Cloudflare R2 configurations. The underlying S3 transport can become reusable, but every provider must pass Mallard's stronger atomic-write and integrity contract.",
        tone="risk",
    )
    add_list(
        doc,
        [
            (
                "Preferred first pilot",
                "Baidu AI Cloud BOS. Its official documentation provides the closest match to Mallard's current protocol: SigV4, both bucket-addressing styles, ListObjectsV2, and conditional PutObject semantics.",
            ),
            (
                "Preferred second pilot",
                "Volcengine TOS. It requires virtual-host addressing and a live test of standard conditional-header and checksum behavior.",
            ),
            (
                "Candidates pending proof",
                "Qiniu Kodo and JD Cloud OSS document strong core S3 and Rust SDK compatibility, but not the exact compare-and-swap contract.",
            ),
            (
                "Top-three provider constraint",
                "Alibaba Cloud OSS, Huawei Cloud OBS, and Tencent Cloud COS do not establish the atomic mutable-head behavior that Mallard requires. Huawei also has a documented ListObjectsV2 gap.",
            ),
        ],
    )
    add_body(
        doc,
        "Recommendation: generalize the R2-only setup into a verified S3-compatible path, then pilot BOS and TOS. Keep all other providers disabled until the same live conformance suite passes. If top-three mainland provider coverage is mandatory, evaluate an append-only publication protocol instead of weakening compare-and-swap safety.",
    )

    doc.add_page_break()
    doc.add_heading("1. Scope and method", level=1)
    doc.add_heading("1.1 Meaning of 'top provider'", level=2)
    add_body(
        doc,
        "No current public report was found for mainland China object-storage market share alone. This assessment therefore uses the current mainland China cloud-infrastructure leaders and adds prominent storage specialists that publish S3 compatibility documentation.",
    )
    add_body(
        doc,
        "Omdia reported Q3 2025 cloud-infrastructure shares of 36% for Alibaba Cloud, 16% for Huawei Cloud, and 9% for Tencent Cloud. The market definition includes compute, storage, databases, and related cloud infrastructure. It is not an object-storage-only ranking. [MK1]",
    )

    doc.add_heading("1.2 Evidence standard", level=2)
    add_list(
        doc,
        [
            "Use official provider API, endpoint, SDK, and release documentation for protocol claims.",
            "Treat marketing statements about S3 compatibility as discovery evidence, not proof of conditional writes or checksum behavior.",
            "State 'not documented' or 'not established' when the reviewed sources do not prove a behavior.",
            "Separate current Mallard product support from lower-level transport feasibility.",
            "Require a credentialed real-bucket test before Mallard enables a provider profile.",
        ],
    )

    doc.add_heading("2. Mallard's current storage contract", level=1)
    doc.add_heading("2.1 Product restrictions", level=2)
    add_body(
        doc,
        "The current frontend accepts only Cloudflare R2 account and bucket URL forms. The backend then requires every S3 entry to use one canonical Cloudflare R2 destination. A manually edited non-R2 record is not a supported workaround because normal validation rejects it. [M1] [M2]",
    )

    doc.add_heading("2.2 Current S3 client behavior", level=2)
    add_list(
        doc,
        [
            "The client uses the AWS Rust S3 SDK with an explicit endpoint.",
            "It uses a static access-key ID and secret-access key. It has no session-token field.",
            "It stores region 'auto' through the R2 setup path.",
            "It always forces path-style bucket addressing.",
            "Every PutObject request includes an AWS S3 SHA-256 checksum.",
            "Every download is also verified against Mallard's expected size and SHA-256 digest. [M3] [M5]",
        ],
    )

    doc.add_heading("2.3 Compatibility gate", level=2)
    add_body(
        doc,
        "Basic upload and download support is not sufficient. A provider must pass all of these checks:",
    )
    add_list(
        doc,
        [
            "Authenticate with AWS SigV4 or a fully compatible request scheme using the configured endpoint and region.",
            "Support the configured path-style or virtual-host bucket-addressing mode.",
            "Support GetObject, HeadObject, PutObject, DeleteObject, and bounded metadata reads.",
            "Support ListObjectsV2 with prefix, StartAfter, continuation tokens, IsTruncated, and stable key traversal.",
            "Create a missing object atomically with If-None-Match: *.",
            "Replace an object atomically with If-Match using its current opaque ETag.",
            "Reject a stale ETag with HTTP 409 or 412.",
            "Return a usable ETag from GET and HEAD responses.",
            "Validate the AWS SHA-256 upload checksum or allow an immediate verified SHA-256 readback before success.",
            "Provide read-after-write and delete visibility that passes Mallard's capability probe.",
        ],
        ordered=True,
    )
    add_body(
        doc,
        "Mallard's existing setup probe already performs conditional create, ETag capture, matching replacement, required stale-ETag rejection, cleanup, and absence verification. The bundle engine uses the same ETag compare-and-swap rule for mutable heads and catalog tags. [M4] [M7]",
    )
    add_callout(
        doc,
        "Safety boundary",
        "Do not emulate compare-and-swap with HeadObject followed by an unconditional PutObject. Another client can update the object between those requests. Do not use ETag as the content checksum. Do not save a provider configuration when the live capability probe fails.",
        tone="risk",
    )

    doc.add_heading("3. Provider compatibility summary", level=1)
    add_body(
        doc,
        "All 'candidate' statuses remain conditional. No mainland provider has passed a live Mallard conformance test in this repository.",
        italic=True,
    )
    provider_rows = [
        (
            "Alibaba Cloud OSS",
            "Virtual only",
            "ListObjectsV2",
            "Standard conditions documented unsupported",
            "Blocked by protocol mismatch",
        ),
        (
            "Huawei Cloud OBS",
            "Virtual only",
            "ListObjects v1; V2 gap",
            "Not established",
            "Blocked by protocol gaps",
        ),
        (
            "Tencent Cloud COS",
            "Virtual for new buckets",
            "V2 needs live proof",
            "Create-only vendor control; update CAS not established",
            "Blocked pending atomic-write proof",
        ),
        (
            "Baidu AI Cloud BOS",
            "Path and virtual",
            "ListObjectsV2",
            "Documented in native PutObject",
            "Preferred first pilot",
        ),
        (
            "Volcengine TOS",
            "Virtual only",
            "ListObjectsV2 on regular buckets",
            "Service support documented; S3 mapping needs proof",
            "Preferred second pilot",
        ),
        (
            "Qiniu Kodo",
            "Path and virtual",
            "ListObjectsV2",
            "Not established",
            "Candidate pending conformance evidence",
        ),
        (
            "JD Cloud OSS",
            "Bucket-as-host documented",
            "ListObjectsV2",
            "Not established",
            "Candidate pending conformance evidence",
        ),
        (
            "China Mobile EOS",
            "Path style in sample",
            "Modern V2 proof absent",
            "Not established",
            "Exploratory only",
        ),
        (
            "China Telecom ZOS",
            "Path and virtual forms documented",
            "V2 proof absent",
            "Not established",
            "Exploratory only",
        ),
        (
            "China Unicom BOSS/OSS",
            "Not established",
            "Not established",
            "Not established",
            "Exploratory only",
        ),
    ]
    table = doc.add_table(rows=1, cols=5)
    for index, value in enumerate(
        ("Provider", "Addressing", "Listing", "Atomic conditional PUT", "Assessment")
    ):
        table.rows[0].cells[index].text = value
    for values in provider_rows[:3]:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    set_table_geometry(table, [1640, 1540, 1540, 2320, 2320])
    style_table(table, body_size=8.3, status_column=4)

    doc.add_page_break()
    continued = doc.add_paragraph()
    continued.paragraph_format.space_before = Pt(0)
    continued.paragraph_format.space_after = Pt(8)
    continued.paragraph_format.keep_with_next = True
    continued_run = continued.add_run("3. Provider compatibility summary - continued")
    set_run_font(continued_run, size=12, color=DARK_BLUE, bold=True)

    table = doc.add_table(rows=1, cols=5)
    for index, value in enumerate(
        ("Provider", "Addressing", "Listing", "Atomic conditional PUT", "Assessment")
    ):
        table.rows[0].cells[index].text = value
    for values in provider_rows[3:]:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    set_table_geometry(table, [1640, 1540, 1540, 2320, 2320])
    style_table(table, body_size=8.3, status_column=4)
    add_table_source(
        doc,
        "Assessment basis: Mallard contract [M1-M7] and official provider sources [A1-CU1]. 'Not established' means that the reviewed public documentation does not prove the required behavior.",
    )

    doc.add_heading("4. Detailed provider findings", level=1)
    doc.add_heading("4.1 Blocked by the current publication protocol", level=2)

    add_provider_heading(
        doc,
        "Alibaba Cloud OSS",
        "Blocked by documented protocol mismatch.",
    )
    add_body(
        doc,
        "OSS supports a documented subset of S3 operations, including core object CRUD and ListObjectsV2. It requires virtual-host addressing. Standard conditional PutObject headers such as If-Match and If-None-Match are explicitly unsupported. OSS provides a vendor create-only control, x-oss-forbid-overwrite, but the reviewed sources do not provide destination ETag compare-and-swap for an existing object. [A1] [A2] [A3]",
    )
    add_body(
        doc,
        "Mallard implication: configurable virtual-host routing would solve only one problem. An adapter could map immutable create to the vendor header on a non-versioned bucket, but it could not safely implement Mallard's documented mutable-head replacement from the reviewed API contract.",
    )

    add_provider_heading(
        doc,
        "Huawei Cloud OBS",
        "Blocked by documented protocol gaps.",
    )
    add_body(
        doc,
        "OBS requires virtual-host addressing for S3 access. The reviewed Huawei documentation states that standard OBS does not support ListObjectsV2. Its PutObject reference does not establish atomic If-Match or If-None-Match writes for destination objects. [H1] [H2] [H3]",
    )
    add_body(
        doc,
        "Mallard implication: a provider adapter would need ListObjects v1 pagination and Huawei-compatible request handling. The publication protocol would still need an atomic-write design that the documented OBS API can implement.",
    )

    add_provider_heading(
        doc,
        "Tencent Cloud COS",
        "Blocked pending atomic-write proof.",
    )
    add_body(
        doc,
        "COS supports AWS SigV4 and S3-style access. Buckets created from 1 January 2024 require virtual-host addressing. COS documents x-cos-forbid-overwrite as a create-only control for non-versioned buckets. The reviewed PutObject documentation does not establish exact destination ETag overwrite compare-and-swap. Tencent also states that COS does not guarantee complete S3 compatibility. [T1] [T2]",
    )
    add_body(
        doc,
        "Mallard implication: core S3 access is plausible after virtual-host support, but the provider must not be enabled until live tests prove ListObjectsV2 and exact stale-ETag rejection. The vendor create-only header is not sufficient for mutable heads.",
    )

    doc.add_heading("4.2 Preferred pilots", level=2)
    add_provider_heading(doc, "Baidu AI Cloud BOS", "Preferred first pilot.")
    add_body(
        doc,
        "BOS documents stock AWS S3 SDK use, AWS SigV4, path-style and virtual-host addressing, core object operations, and ListObjectsV2. Native PutObject documentation explicitly supports If-Match, If-None-Match: *, ETags, and HTTP 412 on a failed precondition. [B1] [B2] [B3] [B4]",
    )
    add_body(
        doc,
        "Required proof: the S3 endpoint must accept the AWS Rust SDK request, enforce all four conditional-write cases, return usable ETags, and either validate the AWS SHA-256 checksum header or pass an immediate SHA-256 readback. Use a mainland HTTPS endpoint and the provider's actual signing region.",
    )

    add_provider_heading(doc, "Volcengine TOS", "Preferred second pilot.")
    add_body(
        doc,
        "TOS provides S3 endpoints but requires virtual-host addressing. A July 2025 release note added If-Match and If-None-Match support to PutObject, CopyObject, and CompleteMultipartUpload. TOS also documents S3 V4 streaming with trailing headers. [V1] [V2]",
    )
    add_body(
        doc,
        "Required proof: the S3 endpoint must map the standard conditional headers to the documented service behavior, reject a stale ETag with 409 or 412, and accept the AWS checksum mode used by Mallard. Use a regular flat/FNS bucket. Do not use an HNS bucket for the first pilot.",
    )

    doc.add_heading("4.3 Candidates pending conformance evidence", level=2)
    add_provider_heading(doc, "Qiniu Kodo", "Candidate pending conformance evidence.")
    add_body(
        doc,
        "Kodo documents path-style and virtual-host addressing, AWS SigV2 and SigV4, core object APIs, and ListObjectsV2. Qiniu also provides an official stock aws-sdk-s3 Rust example. Exact conditional PutObject and modern AWS checksum behavior are not sufficiently documented. [Q1] [Q2] [Q3]",
    )

    add_provider_heading(
        doc,
        "JD Cloud OSS",
        "Candidate pending conformance evidence.",
        page_break_before=True,
    )
    add_body(
        doc,
        "JD Cloud documents SigV4, core S3 object operations, ListObjectsV2, endpoints, and a stock Rust AWS SDK guide. The reviewed PutObject reference does not establish If-Match, If-None-Match, or required 412 behavior. Flexible-checksum behavior also remains unproven. [J1] [J2] [J3] [J4]",
    )

    doc.add_heading("4.4 Exploratory providers", level=2)
    add_provider_heading(doc, "China Mobile Cloud EOS", "Exploratory only.")
    add_body(
        doc,
        "The product page states that EOS provides an S3 interface. The available official sample uses a vendor AWS SDK fork, SigV2, and forced path-style access. The reviewed public evidence does not establish stock Rust AWS SDK SigV4, conditional PutObject, or modern checksum behavior. [C1] [C2]",
    )

    add_provider_heading(doc, "China Telecom Tianyi Cloud ZOS", "Exploratory only.")
    add_body(
        doc,
        "ZOS publishes current S3-style APIs, endpoint forms, and x-amz-content-sha256 usage. The reviewed documentation does not establish the complete Mallard contract for stock AWS SigV4, ListObjectsV2 pagination, conditional writes, and flexible checksums. [CT1] [CT2]",
    )

    add_provider_heading(doc, "China Unicom Cloud BOSS/OSS", "Exploratory only.")
    add_body(
        doc,
        "The product page claims Amazon S3 compatibility, but the reviewed public protocol evidence is insufficient. Endpoint, signing, addressing, ListObjectsV2, conditional writes, and checksum behavior all require provider confirmation and testing. [CU1]",
    )

    doc.add_heading("5. Recommended implementation sequence", level=1)
    add_callout(
        doc,
        "Near-term recommendation",
        "Build one verified generic S3 path and use the same strict capability probe for every provider. Do not add provider names to the UI before the live probe passes.",
        tone="positive",
    )
    add_list(
        doc,
        [
            "Generalize the R2-only setup into an S3-compatible storage profile. Keep R2 account provisioning as a provider-specific flow.",
            "Store the endpoint, bucket, actual signing region, and path-style or virtual-host addressing mode.",
            "Keep static AK/SK support for the first slice. Add a session-token field only if temporary credentials are explicitly in scope.",
            "Make S3 client addressing configurable instead of always forcing path style.",
            "Generalize the existing strict R2 capability probe. Preserve conditional create, ETag capture, matching replacement, stale rejection, checksum verification, deletion, and absence verification.",
            "Add real temporary-bucket integration tests for signing, TLS, DNS, pagination, consistency, checksums, and error status mapping.",
            "Pilot Baidu BOS. Enable it only when the full conformance suite passes.",
            "Add virtual-host support and pilot Volcengine TOS. Then test Qiniu Kodo and JD Cloud with the same suite.",
            "Keep Alibaba OSS, Huawei OBS, and Tencent COS disabled under the current mutable-head publication protocol.",
            "If top-three mainland provider coverage is a product requirement, evaluate append-only immutable head records and deterministic reconciliation as a separate protocol project.",
        ],
        ordered=True,
    )

    doc.add_heading("6. Conformance test checklist", level=1)
    add_list(
        doc,
        [
            (
                "Authentication",
                "Use the exact endpoint, signing region, static credentials, service name, and HTTPS certificate path that Mallard will ship.",
            ),
            (
                "Addressing",
                "Verify both configured routing and bucket naming. Confirm path-style or virtual-host behavior without redirects that invalidate signatures.",
            ),
            (
                "Core object operations",
                "Put, Get, Head, and Delete a zero-byte object and a normal object. Require usable ETags and HTTP 404 for missing reads.",
            ),
            (
                "Atomic create",
                "Create with If-None-Match: *. Repeat the request and require HTTP 409 or 412.",
            ),
            (
                "Atomic replacement",
                "Replace with the current If-Match ETag. Then reuse the stale ETag and require HTTP 409 or 412.",
            ),
            (
                "Integrity",
                "Confirm that the AWS SHA-256 checksum is validated. If it is not supported, require immediate SHA-256 readback before success.",
            ),
            (
                "Listing",
                "Test ListObjectsV2 prefix filtering, StartAfter, continuation tokens, IsTruncated, stable ordering, and pages larger than 1,000 keys.",
            ),
            (
                "Visibility",
                "Verify immediate keyed read after write and HEAD 404 after delete. Treat delayed visibility as a failed provider profile.",
            ),
            (
                "Key behavior",
                "Test Mallard's .mallard/ namespace, Unicode keys, long keys, lexical neighbors, and exact-prefix filtering.",
            ),
            (
                "Failure resolution",
                "Test timeouts, ambiguous write outcomes, stale conditions, missing ETags, and cleanup failure without weakening revision or digest checks.",
            ),
        ],
    )

    doc.add_heading("7. Decision and next actions", level=1)
    table = doc.add_table(rows=1, cols=3)
    for index, value in enumerate(("Decision", "Recommended action", "Acceptance condition")):
        table.rows[0].cells[index].text = value
    decision_rows = [
        (
            "Baidu BOS pilot",
            "Approve",
            "Generic S3 setup and full live conformance suite pass",
        ),
        (
            "Volcengine TOS pilot",
            "Approve after virtual-host support",
            "Standard conditional headers and checksum mode pass",
        ),
        (
            "Qiniu and JD evaluation",
            "Defer until common probe exists",
            "Same provider-neutral suite passes without safety exceptions",
        ),
        (
            "Alibaba, Huawei, Tencent support",
            "Do not enable under current protocol",
            "Documented atomic publication primitive or approved protocol redesign",
        ),
    ]
    for values in decision_rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    set_table_geometry(table, [2440, 2860, 4060])
    style_table(table, body_size=9.0)
    add_table_source(
        doc,
        "Decision basis: current Mallard contract [M1-M7] and provider findings in Sections 3 and 4.",
    )

    doc.add_page_break()
    doc.add_heading("Appendix A - Source register", level=1)
    add_body(
        doc,
        "External protocol claims use official provider sources. The market-context source is an independent analyst report. Mallard references point to the inspected repository state on 1 September 2026.",
        italic=True,
    )

    doc.add_heading("Market context", level=2)
    add_source(
        doc,
        "MK1",
        "Omdia: Mainland China's cloud infrastructure market accelerates to 24% growth in Q3 2025",
        "https://omdia.tech.informa.com/pr/2026/feb/mainland-chinas-cloud-infrastructure-market-accelerates-to-24percent-growth-in-q3-2025",
        "Mainland cloud-infrastructure market shares and market definition.",
    )

    doc.add_heading("Mallard repository references", level=2)
    add_source(
        doc,
        "M1",
        "src/components/project-sync/StorageSettings.tsx:51",
        None,
        "Frontend R2 endpoint parsing and R2-only storage form behavior.",
    )
    add_source(
        doc,
        "M2",
        "crates/mallard-application/src/storage_setup.rs:388",
        None,
        "Backend validation requires every S3 storage to use a canonical Cloudflare R2 destination.",
    )
    add_source(
        doc,
        "M3",
        "crates/mallard-infrastructure/src/storage_core/s3_store.rs:1277",
        None,
        "AWS S3 client construction, static credentials, region, endpoint, and forced path-style addressing.",
    )
    add_source(
        doc,
        "M4",
        "crates/mallard-infrastructure/src/storage_core/s3_store.rs:269",
        None,
        "Capability probe for conditional create, ETag replacement, stale-write rejection, and cleanup.",
    )
    add_source(
        doc,
        "M5",
        "crates/mallard-infrastructure/src/storage_core/s3_store.rs:665",
        None,
        "Conditional PutObject and AWS SHA-256 checksum behavior.",
    )
    add_source(
        doc,
        "M6",
        "crates/mallard-infrastructure/src/storage_core/s3_store.rs:849",
        None,
        "ListObjectsV2 prefix, StartAfter, continuation-token, and pagination behavior.",
    )
    add_source(
        doc,
        "M7",
        "crates/mallard-infrastructure/src/bundle_engine.rs:915",
        None,
        "Immutable object publication and mutable-head compare-and-swap flow.",
    )

    doc.add_heading("Alibaba Cloud OSS", level=2)
    add_source(doc, "A1", "Amazon S3 compatibility", "https://www.alibabacloud.com/help/en/oss/developer-reference/compatibility-with-amazon-s3", "Supported operations, virtual-host requirement, and ETag differences.")
    add_source(doc, "A2", "Conditional PutObject limitation", "https://www.alibabacloud.com/help/en/oss/user-guide/0017-00000245", "Standard PutObject conditional headers are not implemented.")
    add_source(doc, "A3", "PutObject", "https://www.alibabacloud.com/help/en/oss/developer-reference/putobject", "Vendor no-overwrite control and object-write behavior.")

    doc.add_heading("Tencent Cloud COS", level=2)
    add_source(doc, "T1", "S3-compatible application configuration", "https://intl.cloud.tencent.com/document/product/436/34688?lang=en", "Endpoint, signatures, virtual-host requirement for new buckets, and compatibility caveat.")
    add_source(doc, "T2", "PutObject", "https://cloud.tencent.com/document/product/436/7749", "Vendor no-overwrite header and normal overwrite behavior.")

    doc.add_heading("Huawei Cloud OBS", level=2)
    add_source(doc, "H1", "AWS S3 access guidance", "https://support.huaweicloud.com/obs_faq/obs_faq_0274.html", "AWS SigV4 and checksum-mode guidance.")
    add_source(doc, "H2", "ListObjectsV2 limitation", "https://support.huaweicloud.com/intl/en-us/mrs_faq/mrs_03_1102.html", "Standard OBS ListObjectsV2 limitation.")
    add_source(doc, "H3", "PutObject", "https://support.huaweicloud.com/intl/en-us/api-obs/obs_04_0080.html", "OBS PutObject request and response behavior.")

    baidu_sources_heading = doc.add_heading("Baidu AI Cloud BOS", level=2)
    baidu_sources_heading.paragraph_format.page_break_before = True
    add_source(doc, "B1", "S3 endpoints and addressing", "https://cloud.baidu.com/doc/BOS/s/xjwvyq9l4", "Path-style and virtual-host endpoint forms.")
    add_source(doc, "B2", "AWS SDK access guide", "https://cloud.baidu.com/doc/BOS/s/ojwvyq973", "AWS SigV4 and SDK configuration.")
    add_source(doc, "B3", "S3 compatibility matrix", "https://intl.cloud.baidu.com/en/doc/BOS/s/Njwvyq8ho-intl-en", "ListObjectsV2 and core S3 operations.")
    add_source(doc, "B4", "Conditional PutObject", "https://cloud.baidu.com/doc/BOS/s/Ikc5nv3wc", "If-Match, If-None-Match: *, ETag, 412, and native SHA-256 header behavior.")

    doc.add_heading("Volcengine TOS", level=2)
    add_source(doc, "V1", "S3 virtual-host addressing requirement", "https://www.volcengine.com/docs/6349/1288493?lang=zh", "TOS S3 endpoint addressing behavior.")
    add_source(doc, "V2", "Conditional-write release note", "https://www.volcengine.com/docs/6349/79878?lang=en", "July 2025 addition of If-Match and If-None-Match for upload operations.")

    doc.add_heading("Qiniu Kodo", level=2)
    add_source(doc, "Q1", "S3 access domain names", "https://developer.qiniu.com/kodo/4088/s3-access-domainname", "Path-style and virtual-host endpoint forms.")
    add_source(doc, "Q2", "S3 compatibility matrix", "https://developer.qiniu.com/kodo/4087/compatible-s3-api", "Supported core APIs and ListObjectsV2.")
    add_source(doc, "Q3", "AWS Rust SDK example", "https://developer.qiniu.com/kodo/12572/aws-sdk-rust-examples", "Stock aws-sdk-s3 integration example.")

    doc.add_heading("JD Cloud OSS", level=2)
    add_source(doc, "J1", "Regions and endpoints", "https://docs.jdcloud.com/cn/object-storage-service/api/regions-and-endpoints", "Regional S3 endpoint forms.")
    add_source(doc, "J2", "S3 compatibility overview", "https://docs.jdcloud.com/cn/object-storage-service/api/compatibility-api-overview", "Core object APIs and ListObjectsV2 mapping.")
    add_source(doc, "J3", "Rust SDK guide", "https://docs.jdcloud.com/cn/object-storage-service/sdk-rust", "Stock aws-sdk-s3 configuration.")
    add_source(doc, "J4", "PutObject", "https://docs.jdcloud.com/cn/object-storage-service/api/put-object-2", "Published PutObject request behavior; exact conditional headers are not established.")

    doc.add_heading("Other mainland providers", level=2)
    add_source(doc, "C1", "China Mobile Cloud EOS product page", "https://qiye.cmcloud.com/main/zy-eos.html", "Product-level S3 interface statement.")
    add_source(doc, "C2", "China Mobile official Java SDK sample", "https://gitee.com/chinamobile_ecloud/ecloud-eos-sdk-java/blob/master/src/main/java/com/chinamobile/cmss/eos/s3/InitializeSample.java", "Vendor SDK fork, SigV2, endpoint, and forced path-style example.")
    add_source(doc, "CT1", "China Telecom ZOS documentation", "https://www.ctyun.cn/document/10026735/", "Current ZOS object-storage documentation.")
    add_source(doc, "CT2", "China Telecom ZOS PutObject", "https://www.ctyun.cn/document/10026735/11070698", "PutObject and x-amz-content-sha256 use.")
    add_source(doc, "CU1", "China Unicom BOSS product page", "https://www.cucloud.cn/product/BOSS.html", "Product-level Amazon S3 compatibility statement; protocol detail remains insufficient.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build_document()
